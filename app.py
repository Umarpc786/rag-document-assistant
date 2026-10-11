
import io
import os
import re
from typing import List

import pandas as pd
import streamlit as st
from docx import Document
from groq import Groq
from PIL import Image
import pytesseract
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# -----------------------------
# Page setup and readable UI
# -----------------------------
st.set_page_config(
    page_title="DocuMind AI | Document Assistant",
    page_icon="📘",
    layout="wide",
)

st.markdown(
    """
    <style>
    :root { color-scheme: light; }
    html, body, [class*="css"], .stApp { color: #172B4D; }
    .stApp { background-color: #F5F8FC; }
    [data-testid="stHeader"] { background-color: #F5F8FC; }

    [data-testid="stSidebar"] {
        background-color: #EDF3FB;
        border-right: 1px solid #D8E1ED;
    }

    .stMarkdown, .stMarkdown p, .stMarkdown li,
    [data-testid="stWidgetLabel"],
    [data-testid="stCaptionContainer"],
    [data-testid="stExpander"] summary,
    .stRadio label, .stCheckbox label,
    .stSelectbox label, .stTextInput label,
    .stTextArea label, .stFileUploader label,
    .stNumberInput label {
        color: #172B4D !important;
    }

    [data-testid="stCaptionContainer"] {
        color: #526176 !important;
    }

    input, textarea,
    [data-baseweb="select"] > div,
    [data-baseweb="input"] > div,
    [data-baseweb="textarea"] > div {
        background-color: #FFFFFF !important;
        color: #172B4D !important;
        border-color: #B9C7D9 !important;
    }

    input::placeholder, textarea::placeholder {
        color: #66758A !important;
        opacity: 1;
    }

    [data-testid="stFileUploader"] section {
        background-color: #FFFFFF;
        border: 1px dashed #9FB2CA;
        border-radius: 12px;
    }

    [data-testid="stMetric"] {
        background-color: #FFFFFF;
        border: 1px solid #D8E1ED;
        padding: 12px;
        border-radius: 12px;
    }

    [data-testid="stMetricLabel"],
    [data-testid="stMetricValue"] {
        color: #172B4D !important;
    }

    .hero {
        background: linear-gradient(135deg, #174EA6, #2563EB);
        padding: 1.5rem 1.7rem;
        border-radius: 18px;
        margin-bottom: 1rem;
        box-shadow: 0 8px 24px rgba(23, 78, 166, 0.16);
    }

    .hero h1 {
        color: #FFFFFF !important;
        margin: 0;
        font-size: 2rem;
    }

    .hero p {
        color: #F0F6FF !important;
        margin: 0.45rem 0 0;
    }

    div.stButton > button[kind="primary"] {
        background-color: #2563EB;
        border: 1px solid #2563EB;
        color: #FFFFFF !important;
        border-radius: 10px;
        font-weight: 600;
    }

    div.stButton > button[kind="primary"]:hover {
        background-color: #1D4ED8;
        border-color: #1D4ED8;
        color: #FFFFFF !important;
    }

    div.stButton > button[kind="secondary"] {
        background-color: #FFFFFF;
        color: #172B4D !important;
        border: 1px solid #B9C7D9;
        border-radius: 10px;
    }

    .small-note {
        color: #526176 !important;
        font-size: 0.9rem;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="hero">
      <h1>📘 DocuMind AI</h1>
      <p>Your document study assistant — ask questions, simplify notes,
      create quizzes, summarize files, and draft assignments.</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# -----------------------------
# Groq setup
# -----------------------------
def get_api_key() -> str:
    """Read the API key from Streamlit Secrets or an environment variable."""
    try:
        key = st.secrets.get("GROQ_API_KEY", "")
    except Exception:
        key = ""
    return key or os.getenv("GROQ_API_KEY", "")


def call_groq(prompt: str, model: str, temperature: float = 0.2) -> str:
    api_key = get_api_key()

    if not api_key:
        raise ValueError(
            "Groq API key not found. Add GROQ_API_KEY to Streamlit Secrets "
            "or to your local .streamlit/secrets.toml file."
        )

    client = Groq(api_key=api_key)

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a careful document study assistant. Prioritize "
                    "the provided source material. Do not invent facts or claim "
                    "the document says something absent from the supplied context. "
                    "If the context is insufficient, clearly say so. Use clear, "
                    "beginner-friendly English unless asked otherwise."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        temperature=temperature,
        max_tokens=5000,
    )

    return (
        response.choices[0].message.content
        or "The model returned an empty response."
    )


# -----------------------------
# Document text extraction
# -----------------------------
def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {page_number}]\n{text}")

    return "\n\n".join(pages)


def extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_excel(data: bytes, filename: str) -> str:
    file_obj = io.BytesIO(data)
    sheets = pd.read_excel(file_obj, sheet_name=None, header=None)
    parts = []

    for sheet_name, dataframe in sheets.items():
        parts.append(f"Sheet: {sheet_name}")

        if dataframe.empty:
            parts.append("(This sheet is empty.)")
        else:
            dataframe = dataframe.fillna("")
            parts.append(
                dataframe.to_string(
                    index=False,
                    header=False,
                    max_rows=5000,
                )
            )

    return "\n\n".join(parts)


def extract_image(data: bytes) -> str:
    image = Image.open(io.BytesIO(data)).convert("RGB")
    return pytesseract.image_to_string(image)


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        return extract_pdf(data)

    if name.endswith(".docx"):
        return extract_docx(data)

    if name.endswith((".xlsx", ".xls")):
        return extract_excel(data, name)

    if name.endswith((".jpg", ".jpeg", ".png")):
        return extract_image(data)

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="replace")

    if name.endswith(".doc"):
        raise ValueError(
            "Legacy .doc files are not supported directly. Open the file "
            "in Microsoft Word and save it as .docx, then upload it again."
        )

    raise ValueError("Unsupported file type. Please upload a supported document.")


# -----------------------------
# Retrieval: split text into chunks
# -----------------------------
def split_into_chunks(
    text: str,
    chunk_size: int = 900,
    overlap: int = 150,
) -> List[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        if end < len(text):
            boundary = max(
                text.rfind("\n", start, end),
                text.rfind(". ", start, end),
            )
            if boundary > start + chunk_size // 2:
                end = boundary + 1

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(end - overlap, start + 1)

    return chunks


def retrieve_relevant_chunks(
    query: str,
    chunks: List[str],
    top_k: int = 5,
) -> List[str]:
    """Retrieve relevant document sections using local TF-IDF."""
    if not chunks:
        return []

    vectorizer = TfidfVectorizer(
        stop_words="english",
        ngram_range=(1, 2),
    )

    try:
        matrix = vectorizer.fit_transform(chunks + [query])
        query_vector = matrix[-1]
        chunk_matrix = matrix[:-1]

        scores = cosine_similarity(
            query_vector,
            chunk_matrix,
        ).flatten()

        ranked_indices = scores.argsort()[::-1]

        chosen = [
            int(i)
            for i in ranked_indices[:min(top_k, len(chunks))]
            if scores[i] > 0
        ]

        if not chosen:
            return chunks[:min(top_k, len(chunks))]

        return [chunks[i] for i in chosen]

    except ValueError:
        return chunks[:min(top_k, len(chunks))]


def make_context(
    query: str,
    chunks: List[str],
    top_k: int = 6,
) -> str:
    selected = retrieve_relevant_chunks(
        query,
        chunks,
        top_k=top_k,
    )

    return "\n\n--- SOURCE SECTION ---\n\n".join(selected)


def run_task(
    task: str,
    custom_request: str,
    context: str,
    model: str,
) -> str:
    prompt = f"""
USER REQUEST:
{custom_request}

TASK TYPE:
{task}

SOURCE MATERIAL RETRIEVED FROM THE UPLOADED DOCUMENT:
<context>
{context}
</context>

INSTRUCTIONS:
- Use the source material as the main basis for the response.
- If asked for a specific number of questions, produce exactly that number.
- For short questions, keep each question concise. Include answers only if asked.
- For long questions, write descriptive exam-style questions. Include answers only if requested.
- For a summary, cover the main points and keep the wording clear.
- For an assignment, use a suitable title, introduction, organized headings,
  detailed discussion, and conclusion. Use references only when real references
  are supplied or requested. Never invent citations.
- If the document does not provide enough material, say what is missing.
- Format the answer using readable Markdown.
"""

    return call_groq(prompt, model=model)


# -----------------------------
# Sidebar controls
# -----------------------------
with st.sidebar:
    st.header("⚙️ Settings")

    model = st.selectbox(
        "Groq model",
        options=[
            "openai/gpt-oss-120b",
            "openai/gpt-oss-20b",
        ],
        index=0,
        help="Model IDs can change. Check Groq's model list if unavailable.",
    )

    top_k = st.slider(
        "Document sections to retrieve",
        min_value=3,
        max_value=10,
        value=6,
        help="More sections provide more context but can use more tokens.",
    )

    st.markdown("---")
    st.markdown("**Supported formats**")
    st.caption("PDF, DOCX, XLSX, XLS, JPG, JPEG, PNG, TXT")
    st.markdown("---")

    st.markdown(
        '<p class="small-note">Blue + white interface · '
        'Groq-powered generation · Local document retrieval</p>',
        unsafe_allow_html=True,
    )


# -----------------------------
# Upload and process documents
# -----------------------------
st.subheader("1. Upload your document")

uploaded_files = st.file_uploader(
    "Choose one or more files",
    type=[
        "pdf", "docx", "xlsx", "xls",
        "jpg", "jpeg", "png", "txt", "doc",
    ],
    accept_multiple_files=True,
    help="For older .doc files, save them as .docx first.",
)

if "document_text" not in st.session_state:
    st.session_state.document_text = ""

if "chunks" not in st.session_state:
    st.session_state.chunks = []

if "file_names" not in st.session_state:
    st.session_state.file_names = []

if "result" not in st.session_state:
    st.session_state.result = ""


if uploaded_files:
    if st.button("📚 Read and process files", type="primary"):
        all_text = []
        names = []
        errors = []

        with st.spinner("Reading your files..."):
            for file in uploaded_files:
                try:
                    extracted = extract_text(file)

                    if not extracted.strip():
                        errors.append(
                            f"{file.name}: no readable text found. "
                            "Scanned PDFs may need OCR support."
                        )
                        continue

                    all_text.append(
                        f"DOCUMENT: {file.name}\n\n{extracted}"
                    )
                    names.append(file.name)

                except Exception as exc:
                    errors.append(f"{file.name}: {exc}")

        combined_text = "\n\n====================\n\n".join(all_text)
        chunks = split_into_chunks(combined_text)

        st.session_state.document_text = combined_text
        st.session_state.chunks = chunks
        st.session_state.file_names = names
        st.session_state.result = ""

        if names:
            st.success(
                f"Processed {len(names)} file(s) and created "
                f"{len(chunks)} searchable sections."
            )

        for error in errors:
            st.warning(error)


if st.session_state.document_text:
    metric1, metric2, metric3 = st.columns(3)

    metric1.metric("Files read", len(st.session_state.file_names))
    metric2.metric(
        "Text characters",
        f"{len(st.session_state.document_text):,}",
    )
    metric3.metric(
        "Searchable sections",
        len(st.session_state.chunks),
    )

    with st.expander("View extracted text"):
        st.text_area(
            "Extracted document content",
            value=st.session_state.document_text[:30000],
            height=250,
            disabled=True,
        )

        if len(st.session_state.document_text) > 30000:
            st.caption(
                "Preview limited to 30,000 characters. "
                "The full text remains available to the app."
            )

    st.subheader("2. What would you like the assistant to do?")

    task = st.selectbox(
        "Choose a task",
        [
            "Ask a question",
            "Make short questions",
            "Make long questions",
            "Make MCQs",
            "Summarize document",
            "Explain in easy wording",
            "Create an assignment",
            "Custom request",
        ],
    )

    col1, col2 = st.columns([2, 1])

    with col1:
        if task == "Ask a question":
            request = st.text_area(
                "Your question",
                placeholder=(
                    "Example: Explain the difference between "
                    "supervised and unsupervised learning."
                ),
            )

        elif task == "Make short questions":
            count = st.number_input(
                "Number of short questions",
                min_value=1,
                max_value=50,
                value=10,
            )
            include_answers = st.checkbox(
                "Include answers",
                value=False,
            )
            request = (
                f"Create exactly {count} short exam questions. "
                + (
                    "Include a short answer after each question."
                    if include_answers
                    else "Do not include answers."
                )
            )

        elif task == "Make long questions":
            count = st.number_input(
                "Number of long questions",
                min_value=1,
                max_value=30,
                value=10,
            )
            include_answers = st.checkbox(
                "Include detailed answers",
                value=False,
            )
            request = (
                f"Create exactly {count} long/descriptive exam questions. "
                + (
                    "Include detailed answers."
                    if include_answers
                    else "Do not include answers."
                )
            )

        elif task == "Make MCQs":
            count = st.number_input(
                "Number of MCQs",
                min_value=1,
                max_value=50,
                value=10,
            )
            include_answers = st.checkbox(
                "Show correct answers",
                value=True,
            )
            request = (
                f"Create exactly {count} multiple-choice questions, "
                "each with four options (A, B, C, D). "
                + (
                    "Mark the correct answer and briefly explain it."
                    if include_answers
                    else "Do not reveal the correct answers."
                )
            )

        elif task == "Summarize document":
            request = st.text_area(
                "Summary instructions (optional)",
                placeholder="Example: Summarize in easy English using headings.",
            )
            if not request.strip():
                request = (
                    "Summarize the uploaded document. Cover the key ideas "
                    "and use clear headings."
                )

        elif task == "Explain in easy wording":
            request = st.text_area(
                "What should be explained?",
                placeholder="Example: Explain the whole document in easy English.",
            )
            if not request.strip():
                request = (
                    "Explain the main ideas in the uploaded document "
                    "in easy English."
                )

        elif task == "Create an assignment":
            topic = st.text_input(
                "Assignment topic",
                placeholder="Example: Applications of Artificial Neural Networks",
            )
            length = st.selectbox(
                "Assignment length",
                [
                    "Short (about 500 words)",
                    "Medium (about 1000 words)",
                    "Long (about 1500 words)",
                ],
                index=1,
            )
            style = st.selectbox(
                "Writing style",
                [
                    "Simple student English",
                    "Formal academic English",
                ],
                index=0,
            )
            extra = st.text_area(
                "Extra instructions (optional)",
                placeholder="Example: Include 5 headings and practical examples.",
            )
            request = (
                f"Write an assignment about: {topic}. "
                f"Length: {length}. Style: {style}. "
                f"Extra instructions: {extra or 'None'}."
            )

        else:
            request = st.text_area(
                "Write your own request",
                placeholder="Example: Create revision notes for Chapter 3.",
            )

    with col2:
        st.markdown("#### Quick tips")
        st.caption("• Ask one clear task at a time.")
        st.caption("• For assignments, enter a specific topic.")
        st.caption("• Responses use retrieved document sections.")
        st.caption("• Long files may need more than one request.")

    if st.button("✨ Generate result", type="primary"):
        if not request.strip():
            st.error("Please enter a question, topic, or request.")
        else:
            try:
                context = make_context(
                    request,
                    st.session_state.chunks,
                    top_k=top_k,
                )

                if not context.strip():
                    st.error(
                        "No relevant text was found in the uploaded document."
                    )
                else:
                    with st.spinner("Groq is preparing your response..."):
                        answer = run_task(
                            task,
                            request,
                            context,
                            model,
                        )

                    st.session_state.result = answer

            except Exception as exc:
                st.error(f"Could not generate a response: {exc}")

    if st.session_state.result:
        st.subheader("3. Your result")
        st.markdown(st.session_state.result)

        st.download_button(
            "⬇️ Download result as TXT",
            data=st.session_state.result.encode("utf-8"),
            file_name="document_assistant_result.txt",
            mime="text/plain",
        )

else:
    st.info(
        "Upload a file and click **Read and process files** to begin."
    )
    st.markdown(
        """
        **Try these example requests after uploading a document:**
        - Make 10 short questions from this file.
        - Make 10 long questions with answers.
        - Summarize the document in easy English.
        - Explain Chapter 1 as if I am a beginner.
        - Create an assignment on [your topic] using this material.
        """
    )

st.markdown("---")
st.caption(
    "DocuMind AI · Built with Streamlit and Groq. "
    "Check generated work against your source before submitting it."
)
